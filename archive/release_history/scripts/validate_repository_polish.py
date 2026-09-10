"""Static, read-only validation: no model imports, pickle loads or fitting.

Usage: python -B scripts/validate_repository_polish.py --out NEW_DIRECTORY
Receipts are created exclusively. Existing artifacts are never written.
"""
from __future__ import annotations

import argparse
import ast
import csv
import importlib.metadata
import io
import json
from pathlib import Path
import pickletools
import re
import shutil
import subprocess
from urllib.parse import unquote
import xml.etree.ElementTree as ET
import zipfile

from repository_polish_audit import ROOT, sha256

PACKAGE = ROOT / 'FINAL_SUBMISSION'
SOURCE = ROOT / 'candidate/final_submission'


def validate() -> tuple[dict, str]:
    """Return explicit checks and limitations without scientific execution."""
    try:
        import nbformat
        notebook_schema_error = None
    except Exception as exc:
        nbformat = None
        notebook_schema_error = f'{type(exc).__name__}: {exc}'
    import numpy as np
    from PIL import Image
    from pptx import Presentation
    from packaging.requirements import Requirement

    checks = []
    details = {}
    if notebook_schema_error:
        checks.append({'name': 'nbformat_import_environment', 'status': 'BLOCKED', 'detail': notebook_schema_error})

    def check(name: str, ok: bool, detail: object = '') -> None:
        checks.append({'name': name, 'status': 'PASS' if ok else 'FAIL', 'detail': detail})

    manifest = json.loads((PACKAGE / 'manifest.json').read_text())
    expected = set(manifest['files']) | {'manifest.json', 'checksums.sha256'}
    actual = {p.relative_to(PACKAGE).as_posix() for p in PACKAGE.rglob('*') if p.is_file()}
    check('package_membership', expected == actual, {'missing': sorted(expected-actual), 'extra': sorted(actual-expected)})
    bad = [n for n, entry in manifest['files'].items() if not (PACKAGE/n).is_file() or sha256(PACKAGE/n) != entry['sha256'] or (PACKAGE/n).stat().st_size != entry['bytes']]
    check('package_manifest_hashes', not bad, bad)
    sums = {}
    for line in (PACKAGE/'checksums.sha256').read_text().splitlines():
        digest, name = line.split('  ', 1)
        sums[name] = digest
    check('checksum_coverage', set(sums) == actual-{'checksums.sha256'}, len(sums))
    bad = [n for n, digest in sums.items() if not (PACKAGE/n).is_file() or sha256(PACKAGE/n) != digest]
    check('package_checksums', not bad, bad)
    bad = [n for n, entry in manifest['files'].items() if entry['source'].startswith('candidate/') and sha256(ROOT/entry['source']) != sha256(PACKAGE/n)]
    check('canonical_copies_byte_identical', not bad, bad)
    allowed = {'README.md', 'presentation.pptx', 'presentation.pdf', 'written_explanation', 'solution_notebook.ipynb', 'requirements.txt', 'reproduce.py', 'model', 'results_summary', 'manifest.json', 'checksums.sha256'}
    check('top_level_allowlist', {p.name for p in PACKAGE.iterdir()} == allowed)
    inference_manifest = json.loads((PACKAGE/'model/INFERENCE_MANIFEST.json').read_text())
    check('frozen_model_hashes', all(sha256(PACKAGE/'model'/n) == h for n, h in inference_manifest['files'].items()), list(inference_manifest['files']))
    check('feature_contract', len(inference_manifest['feature_names']) == 60 and inference_manifest['pool_rows'] == 25992)
    lock = json.loads((PACKAGE/'model/SELECTION_LOCK.json').read_text())
    check('selected_model_unchanged', lock['candidate'] == 'Coordinate_RF')

    opened = []
    opening_errors = []
    # Open every file in the new distribution, without deserializing the pickle.
    for p in sorted(PACKAGE.rglob('*')):
        if not p.is_file():
            continue
        try:
            if p.suffix == '.json':
                json.loads(p.read_text(encoding='utf-8'))
            elif p.suffix == '.ipynb':
                if nbformat is not None:
                    nbformat.validate(nbformat.read(p, as_version=4))
                else:
                    notebook = json.loads(p.read_text(encoding='utf-8'))
                    assert notebook['nbformat'] == 4 and isinstance(notebook['cells'], list)
                    assert len({c['id'] for c in notebook['cells']}) == len(notebook['cells'])
                    for cell in notebook['cells']:
                        assert cell['cell_type'] in {'code', 'markdown', 'raw'}
                        assert isinstance(cell['source'], (str, list))
                        assert isinstance(cell['metadata'], dict)
                        if cell['cell_type'] == 'code':
                            assert isinstance(cell['outputs'], list)
            elif p.suffix == '.py':
                ast.parse(p.read_text(encoding='utf-8'))
            elif p.suffix == '.npz':
                with zipfile.ZipFile(p) as z:
                    assert z.testzip() is None
                with np.load(p, allow_pickle=False) as arrays:
                    for key in arrays.files:
                        arrays[key]
            elif p.suffix == '.pkl':
                # Parse opcodes only: this does not invoke reducers or import classes.
                last = None
                for opcode, argument, position in pickletools.genops(p.read_bytes()):
                    last = opcode.name
                assert last == 'STOP'
            elif p.suffix == '.pptx':
                with zipfile.ZipFile(p) as z:
                    assert z.testzip() is None
                    for name in z.namelist():
                        if name.endswith(('.xml', '.rels')):
                            ET.fromstring(z.read(name))
                        elif name.startswith('ppt/media/') and name.endswith(('.png', '.jpg', '.jpeg')):
                            with Image.open(io.BytesIO(z.read(name))) as im:
                                im.verify()
                deck = Presentation(p)
                details['slides'] = len(deck.slides)
            elif p.suffix == '.pdf':
                payload = p.read_bytes()
                assert payload.startswith(b'%PDF-') and b'%%EOF' in payload[-2048:]
            elif p.suffix == '.csv':
                with p.open(encoding='utf-8', newline='') as stream:
                    rows = list(csv.reader(stream))
                    assert rows and all(len(row) == len(rows[0]) for row in rows)
            else:
                p.read_text(encoding='utf-8')
            opened.append(p.relative_to(PACKAGE).as_posix())
        except Exception as exc:
            opening_errors.append({'file': p.relative_to(PACKAGE).as_posix(), 'error': str(exc)})
    check('all_package_files_structurally_open', not opening_errors, {'opened': len(opened), 'errors': opening_errors, 'pickle': 'opcode syntax only, NOT deserialized'})

    original = json.loads((SOURCE/'notebook/FINAL_NOTEBOOK.ipynb').read_text(encoding='utf-8'))
    derived = json.loads((PACKAGE/'solution_notebook.ipynb').read_text(encoding='utf-8'))
    # Validate the exact bundled 4.5 schema without importing the contaminated
    # jsonschema resource registry. This is independent of the UI import failure.
    import importlib.util
    import fastjsonschema
    schema_path = Path(importlib.util.find_spec('nbformat').origin).parent / 'v4/nbformat.v4.5.schema.json'
    fastjsonschema.compile(json.loads(schema_path.read_text(encoding='utf-8')))(derived)
    check('full_notebook_v4_5_schema', True, {'schema_sha256': sha256(schema_path), 'engine': 'fastjsonschema; no model execution'})
    for cell in original['cells']:
        source = cell['source']
        is_list = isinstance(source, list)
        text = ''.join(source) if is_list else source
        text = text.replace("ROOT/'evidence/", "ROOT/'model/evidence/").replace("ROOT/'tables/", "ROOT/'model/tables/")
        if cell['cell_type'] == 'markdown':
            for name in ('SOURCE_REFIT.json', 'PREFLIGHT.json', 'SHA256SUMS.txt'):
                text = text.replace('evidence/'+name, 'model/evidence/'+name)
            text = text.replace('reports/METHODS.md', 'model/reports/METHODS.md')
        cell['source'] = text.splitlines(keepends=True) if is_list else text
    check('notebook_exact_path_only_derivation', original == derived)
    missing_paths = []
    for cell in derived['cells']:
        text = ''.join(cell['source']) if isinstance(cell['source'], list) else cell['source']
        if cell['cell_type'] == 'code':
            ast.parse(text)
            for path in re.findall(r"ROOT\s*/\s*['\"]([^'\"]+)['\"]", text):
                if not (PACKAGE/path).exists():
                    missing_paths.append(path)
    check('notebook_literal_runtime_paths', not missing_paths, missing_paths)
    # Verify support/query schemas without calling the adaptation function.
    reference_summary = {}
    for budget in [5, 25, 50, 100, 200]:
        with np.load(PACKAGE/f'model/verification_inputs/b{budget}.npz', allow_pickle=False) as a:
            s, y, q, pred = (a[n] for n in ('support', 'support_labels', 'query', 'expected'))
            ok = (s.dtype.kind in 'iu' and q.dtype.kind in 'iu' and len(np.unique(s)) == len(s)
                  and len(np.unique(q)) == len(q) and not np.intersect1d(s, q).size
                  and np.all((s >= 0) & (s < 25992)) and np.all((q >= 0) & (q < 25992))
                  and len(s) == budget*4 and all(np.sum(y == c) == budget for c in [1, 2, 3, 4])
                  and len(pred) == len(q) and set(np.unique(pred)) <= {1, 2, 3, 4})
            reference_summary[str(budget)] = {'support': len(s), 'query': len(q), 'valid': bool(ok)}
    check('five_reference_episode_schemas', all(v['valid'] for v in reference_summary.values()), reference_summary)

    audit = json.loads((SOURCE/'evidence/final_complete.json').read_text())
    arithmetic_errors = []
    grouped = {}
    for i, record in enumerate(audit['records']):
        cm = np.asarray(record['confusion'])
        denominator = cm.sum(0) + cm.sum(1)
        f1 = np.divide(2*np.diag(cm), denominator, out=np.zeros(4), where=denominator != 0).mean()
        if not np.isclose(f1, record['f1'], atol=1e-12, rtol=0):
            arithmetic_errors.append(i)
        grouped.setdefault((str(record['budget']), record['method']), []).append(record['f1'])
    check('locked_confusion_f1_arithmetic', not arithmetic_errors, {'records': len(audit['records']), 'errors': arithmetic_errors})
    errors = []
    for (budget, method), values in grouped.items():
        row = audit['summary'][budget][method]
        if not (len(values) == row['episodes'] and np.isclose(np.mean(values), row['mean'], atol=1e-12, rtol=0) and np.isclose(np.std(values, ddof=0), row['population_sd'], atol=1e-12, rtol=0)):
            errors.append([budget, method])
    check('locked_summary_arithmetic', not errors, errors)
    table_errors = []
    for folder in (SOURCE/'tables', ROOT/'candidate/tables', PACKAGE/'results_summary'):
        with (folder/'learning_curve.csv').open(encoding='utf-8', newline='') as stream:
            for row in csv.DictReader(stream):
                ref = audit['summary'][row['budget']][row['method']]
                if not all(np.isclose(float(row[k]), ref[k], atol=1e-12, rtol=0) for k in ('mean', 'population_sd', 'episodes')):
                    table_errors.append([str(folder), row['budget'], row['method']])
    check('learning_curve_tables_match_locked_summary', not table_errors, table_errors)

    deck = Presentation(PACKAGE/'presentation.pptx')
    slide_text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    off_slide = [{'slide': i+1, 'shape': shape.name} for i, slide in enumerate(deck.slides) for shape in slide.shapes if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width or shape.top + shape.height > deck.slide_height]
    check('slide_shapes_within_canvas', not off_slide, off_slide)
    import hashlib
    figure_sources = {}
    for folder in ('candidate/final_submission/figures', 'candidate/figures/publication', 'candidate/presentation'):
        for path in (ROOT/folder).rglob('*.png'):
            figure_sources.setdefault(sha256(path), []).append(path.relative_to(ROOT).as_posix())
    with zipfile.ZipFile(PACKAGE/'presentation.pptx') as archive:
        media = {name: figure_sources.get(hashlib.sha256(archive.read(name)).hexdigest(), []) for name in archive.namelist() if name.startswith('ppt/media/') and name.endswith('.png')}
    check('embedded_slide_figures_match_existing_assets', bool(media) and all(media.values()), media)
    pdf_command = shutil.which('pdftotext')
    pdf_text = ''
    if pdf_command:
        result = subprocess.run([pdf_command, '-layout', str(PACKAGE/'presentation.pdf'), '-'], capture_output=True)
        pdf_text = result.stdout.decode('utf-8', errors='replace')
        check('pdf_text_extraction', result.returncode == 0 and len(pdf_text) > 1000, {'pages': pdf_text.count('\f'), 'stderr': result.stderr.decode(errors='replace')})
        check('pdf_slide_page_count', pdf_text.count('\f') == len(deck.slides))
    else:
        checks.append({'name': 'pdf_text_extraction', 'status': 'NOT_RUN', 'detail': 'pdftotext unavailable'})
    texts = {
        'package_readme': (PACKAGE/'README.md').read_text(encoding='utf-8'),
        'executive_summary': (ROOT/'PROJECT_EXECUTIVE_SUMMARY.md').read_text(encoding='utf-8'),
        'method_guide': (ROOT/'docs/METHOD.md').read_text(encoding='utf-8'),
        'locked_report': (SOURCE/'reports/FINAL_REPORT.md').read_text(encoding='utf-8'),
        'notebook': json.dumps(derived, ensure_ascii=False),
        'pptx': slide_text, 'pdf': pdf_text,
    }
    missing_numbers = {}
    numbers = [f"{audit['summary'][str(b)]['Coordinate_RF'][field]:.6f}" for b in [5,25,50,100,200] for field in ('mean','population_sd')]
    for name, text in texts.items():
        absent = [n for n in numbers if n not in text]
        if absent:
            missing_numbers[name] = absent
    check('all_five_mean_sd_tokens_in_current_materials', not missing_numbers, missing_numbers)
    explanation = (PACKAGE/'written_explanation/written_justification.txt').read_text(encoding='utf-8')
    check('explanation_five_means_match', all(f"{audit['summary'][str(b)]['Coordinate_RF']['mean']:.6f}" in explanation for b in [5,25,50,100,200]))
    normalize = lambda text: re.sub(r'\s+', '', text)
    explanation_slide_text = '\n'.join(shape.text for slide in list(deck.slides)[-2:] for shape in slide.shapes if shape.has_text_frame and not shape.text.startswith('Written explanation ('))
    check('explanation_embedded_in_slides', normalize(explanation) == normalize(explanation_slide_text))
    details['explanation_words_whitespace'] = len(explanation.split())
    details['abstract_words_whitespace'] = len(' '.join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame).split())
    details['word_limit_status'] = '340 words: exceeds conservative 300-word interpretation; below 500. Organizer confirmation required.'

    # Link only explicit Markdown links, not prose/backtick paths or external URLs.
    baseline = json.loads((ROOT/'reproducibility/repository_polish/baseline.json').read_text())['files']
    legacy_links, new_links = [], []
    documents = [ROOT/n for n in baseline if n.endswith('.md')]
    documents += [p for p in ROOT.glob('*.md') if p.relative_to(ROOT).as_posix() not in baseline]
    documents += [p for folder in ('docs', 'deliverables', 'FINAL_SUBMISSION') for p in (ROOT/folder).rglob('*.md') if p.relative_to(ROOT).as_posix() not in baseline]
    for path in documents:
        text = path.read_text(encoding='utf-8', errors='replace')
        text = re.sub(r'```.*?```', '', text, flags=re.S)
        for target in re.findall(r'(?<!!)\[[^\]]+\]\(([^)]+)\)|!\[[^\]]*\]\(([^)]+)\)', text):
            raw = next(t for t in target if t).strip().split(' "')[0].strip('<>')
            if re.match(r'^[a-zA-Z][\w+.-]*:', raw) or raw.startswith('#'):
                continue
            relative = unquote(raw.split('#')[0])
            if relative and not (path.parent/relative).exists():
                item = {'file': path.relative_to(ROOT).as_posix(), 'target': raw}
                (legacy_links if item['file'] in baseline else new_links).append(item)
    check('new_explicit_local_markdown_links', not new_links, new_links)
    details['legacy_broken_links'] = legacy_links
    details['link_scope'] = 'Explicit inline local Markdown file targets only; no heading anchors, external URLs, prose or notebook prose references certified.'

    installed = {}
    pins_ok = True
    for line in (PACKAGE/'requirements.txt').read_text().splitlines():
        if not line.strip() or line.startswith('#'):
            continue
        requirement = Requirement(line)
        try:
            version = importlib.metadata.version(requirement.name)
            match = requirement.specifier.contains(version)
        except importlib.metadata.PackageNotFoundError:
            version, match = None, False
        installed[requirement.name] = {'required': str(requirement.specifier), 'installed': version, 'matches': match}
        pins_ok &= match
    check('installed_package_pins', pins_ok, installed)
    details['environment'] = {'python': __import__('sys').version, 'clean_install_tested': False}
    # Check new schematics structurally; scientific figures are never regenerated.
    diagram_errors = []
    for p in (ROOT/'figures/repository_polish').glob('*'):
        try:
            if p.suffix == '.svg':
                ET.parse(p)
            elif p.suffix == '.png':
                with Image.open(p) as im:
                    im.verify()
        except Exception as exc:
            diagram_errors.append([p.name, str(exc)])
    check('new_diagram_formats', not diagram_errors, diagram_errors)
    changed, missing = [], []
    for name, entry in baseline.items():
        path = ROOT/name
        if not path.is_file():
            missing.append(name)
        elif sha256(path) != entry['sha256']:
            changed.append(name)
    check('all_original_files_preserved', not changed and not missing, {'inventoried': len(baseline), 'changed': changed, 'missing': missing})
    scientific_changes = [name for name in changed + missing if not name.endswith('.onetoc2')]
    check('all_original_non_onenote_files_preserved', not scientific_changes, scientific_changes)
    historical = json.loads((ROOT/'SHA256SUMS.json').read_text())['files']
    historical_bad = [name for name, digest in historical.items() if not (ROOT/name).is_file() or sha256(ROOT/name) != digest]
    details['historical_root_manifest_mismatches'] = historical_bad
    details['historical_manifest_note'] = 'Reported separately from this pass preservation. Never rewrite a frozen manifest to make it pass.'
    details['runtime_not_run'] = ['source training', 'support RF adaptation', 'prediction replay', 'notebook execution', 'pickle deserialization', 'fresh PowerPoint rendering', 'clean environment installation']
    visual_receipt = ROOT/'reproducibility/repository_polish/visual_review/receipt.json'
    details['separate_pdf_render_receipt'] = json.loads(visual_receipt.read_text(encoding='utf-8')) if visual_receipt.exists() else 'NOT_RUN'
    return {'checks': checks, 'details': details, 'failed_checks': sum(c['status']=='FAIL' for c in checks), 'scope': 'Read-only/static and saved-evidence arithmetic; not scientific reproduction'}, pdf_text


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=False)
    payload, pdf_text = validate()
    with (args.out/'validation.json').open('x', encoding='utf-8') as stream:
        json.dump(payload, stream, indent=2)
    with (args.out/'presentation_text.txt').open('x', encoding='utf-8') as stream:
        stream.write(pdf_text)
    print(json.dumps({'checks': len(payload['checks']), 'failed': payload['failed_checks'], 'legacy_broken_links': len(payload['details']['legacy_broken_links']), 'historical_root_hash_mismatches': len(payload['details']['historical_root_manifest_mismatches'])}, indent=2))
    raise SystemExit(1 if payload['failed_checks'] else 0)


if __name__ == '__main__':
    main()
