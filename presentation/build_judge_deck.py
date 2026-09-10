"""Build ten presentation-only slides from frozen tables; never import model code.

Run: python -B presentation/build_judge_deck.py
PDF/PNG export: powershell.exe -NoProfile -File presentation/export_judge_deck.ps1
Requires the already available python-pptx, Pillow and Matplotlib packages.
"""
from pathlib import Path
import csv
import hashlib
import json
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'presentation'
TABLE = ROOT / 'FINAL_SUBMISSION/results_summary/learning_curve.csv'
ROWS = list(csv.DictReader(TABLE.open(encoding='utf-8', newline='')))
VALUES = {int(r['budget']): r for r in ROWS if r['method'] == 'Coordinate_RF'}
assert f"{float(VALUES[200]['mean']):.6f}" == '0.749299'
NAVY, TEAL, BLUE, MUTED, WASH, ORANGE = '142B40', '008E78', '2879B9', '526478', 'EFF4F8', 'AE5B17'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

def text(s, value, x, y, w, h, size=23, color=NAVY, bold=False):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(10)
    return shape

def panel(s, x, y, w, h, color=WASH):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    return sh

def slide(title, section, source, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(TEAL); bar.line.fill.background()
    text(s, section.upper(), .6, .32, 12, .3, 12, TEAL, True)
    text(s, title, .6, .9, 12.1, .95, 32, bold=True)
    text(s, source, .6, 7.02, 11.6, .26, 10, MUTED)
    text(s, f'{len(prs.slides):02d} / 10', 12.05, 7.02, .8, .26, 10, MUTED)
    s.notes_slide.notes_text_frame.text = notes
    return s

def image(s, p, x, y, w, h):
    with Image.open(p) as im:
        ratio = min(w / im.width, h / im.height)
        iw, ih = im.width * ratio, im.height * ratio
    s.shapes.add_picture(str(p), Inches(x + (w-iw)/2), Inches(y+(h-ih)/2), width=Inches(iw), height=Inches(ih))

def card(s, x, y, w, heading, body, h=2.4):
    panel(s, x, y, w, h)
    text(s, heading, x+.22, y+.22, w-.44, .65, 25, TEAL, True)
    text(s, body, x+.22, y+1, w-.44, h-1.1, 21)

# A new presentation view only; all means/SD come directly from the immutable CSV.
fig, ax = plt.subplots(figsize=(10.8, 5.25), dpi=200)
for method, label, color in [('EXP010','EXP-010','#838C96'), ('EXPF','EXP-F','#AE5B17'),
                              ('ASTRA_AGF','ASTRA','#2879B9'), ('Coordinate_RF','Coordinate-RF','#008E78')]:
    rs = sorted((r for r in ROWS if r['method'] == method), key=lambda r: int(r['budget']))
    ax.errorbar([math.log2(int(r['budget'])) for r in rs], [float(r['mean']) for r in rs],
                yerr=[float(r['population_sd']) for r in rs], label=label, color=color,
                marker='o', linewidth=2.3, capsize=4)
ax.set_xticks([math.log2(b) for b in VALUES], [str(b) for b in VALUES])
ax.set_xlabel('Labels per class (log2 spacing)', fontsize=13)
ax.set_ylabel('Macro-F1 ± population SD', fontsize=13)
ax.set_ylim(.49, .78); ax.grid(axis='y', alpha=.2)
ax.spines[['top','right']].set_visible(False)
ax.legend(loc='lower right', frameon=False, fontsize=12)
fig.tight_layout(); fig.savefig(OUT/'judge_learning_curve.png', facecolor='white'); plt.close(fig)

s = slide('Building-age transfer with scarce labels', 'StatsGeeks · Coordinate-RF',
          'Frozen evidence: FINAL_SUBMISSION/results_summary/learning_curve.csv',
          'Speaker 1. First-slide abstract, under 150 words. Introduce the problem and the exact scope, not just the score.')
text(s, 'Madrid → Amsterdam', .7, 1.95, 7, .7, 35, TEAL, True)
text(s, 'We classify four building-construction eras from 30 m Landsat pixels. '
        'Coordinate-RF combines a Madrid-learned prior with a target-support random forest, '
        'whitened spectral features and spatial context. On the frozen Amsterdam random-pixel '
        'protocol it reaches 0.749299 macro-F1 at 200 labels per class. '
        'The method uses the full unlabelled target pool and an additional development label bank; '
        'spatial extrapolation remains a limitation.', .7, 2.9, 8.05, 3.65, 23)
card(s, 9.1, 2.1, 3.5, '0.749299', 'Locked macro-F1\n200 labels/class\n200 paired episodes', 3.5)

s = slide('A useful urban signal, with imperfect labels', '01 · Problem & motivation',
          'Data: organizer notebooks; docs/MASTER_PLAN.md dataset summary',
          'Speaker 1. Age supports retrofit and energy planning. Explain mixed-pixel labels and different oldest-era boundaries.')
card(s, .7, 2, 3.8, 'Why age?', 'Retrofit planning\nEnergy-policy evidence\nIncomplete records', 3.5)
card(s, 4.77, 2, 3.8, 'What data?', 'Six spectral bands\n1984–2025 imagery\n30 m mixed pixels', 3.5)
card(s, 8.84, 2, 3.8, 'What task?', 'Four construction eras\nMadrid source city\nAmsterdam target city', 3.5)
text(s, '76,263 source pixels  ·  25,992 target pixels  ·  60 spectral features', .75, 6, 11.9, .55, 22, MUTED)

s = slide('Transfer must work when labels are scarce', '02 · Challenge understanding',
          'Protocol: candidate/reports/METHODS.md; locked learning-curve CSV',
          'Speaker 1 → Speaker 2. The 25-shot result is important to the rubric. Macro-F1 gives each class equal weight.')
card(s, .7, 2, 5.8, 'Cross-city shift', 'Spectral and spatial patterns differ.\nSource knowledge needs adaptation.', 2.5)
card(s, 6.8, 2, 5.8, 'Few-label adaptation', '5 / 25 / 50 / 100 / 200 labels per class.\nOnly designated support labels fit the local RF.', 2.5)
text(s, '25 labels/class = 100 target support labels', .85, 5.05, 11.6, .6, 28, TEAL, True)
text(s, 'The additional 800-label development bank is separate research supervision.', .85, 6, 11.6, .65, 21, MUTED)

s = slide('One source prior, one local adaptation branch', '03 · Pipeline',
          'Method: candidate/reports/METHODS.md; frozen model configuration',
          'Speaker 2. Historical learning pipeline, not executed during release. Two source RFs produce the cached prior. Only support labels fit the target RF.')
card(s, .7, 1.95, 5.75, 'Madrid source branch', 'Contextual spectral features\nCovariance alignment + refinement\nFrozen source probability prior', 2.7)
card(s, 6.85, 1.95, 5.75, 'Amsterdam local branch', 'Full-pool adaptive whitening\nTwo standardized coordinates\nSupport-only 200-tree RF', 2.7)
text(s, '↓', 3.2, 4.73, 1, .5, 30, TEAL, True)
text(s, '↓', 9.35, 4.73, 1, .5, 30, TEAL, True)
panel(s, 1.7, 5.4, 9.95, .9, NAVY)
text(s, 'Blend probabilities  →  nine-neighbour smoothing  →  classes', 1.95, 5.58, 9.5, .55, 23, 'FFFFFF', True)

s = slide('The contribution is the combination', '04 · Design rationale & creativity',
          'Mechanism evidence: candidate/reports/RESULTS.md and METHODS.md',
          'Speaker 2. Explain rationale without presenting entangled ablations as additive causal effects. No novelty claim for RF itself.')
card(s, .7, 2, 3.8, 'Source reuse', 'Source probability weight:\n20 / (20 + 4b)\nLess weight as support grows.', 3.45)
card(s, 4.77, 2, 3.8, 'Local geometry', 'Adaptive whitening\nSpectral + coordinate inputs\nBalanced random forest', 3.45)
card(s, 8.84, 2, 3.8, 'Spatial context', 'Label-free smoothing\nNine spatial neighbours\nFull-pool access required', 3.45)
text(s, 'b = labels/class. Components interact; ablation bars are not additive causal effects.', .75, 6, 11.9, .65, 20, MUTED)

s = slide('Performance improves as target support grows', '05 · Locked learning curve',
          'Source: FINAL_SUBMISSION/results_summary/learning_curve.csv · no new predictions',
          'Speaker 3. Explain every budget and population SD. These are 200 paired episodes per budget on one fixed city. ASTRA, EXP-F and EXP-010 are historical comparison methods.')
image(s, OUT/'judge_learning_curve.png', .65, 1.75, 9.15, 4.8)
text(s, '25 / class', 10.05, 2.1, 2.6, .45, 20, MUTED)
text(s, f"{float(VALUES[25]['mean']):.6f}", 10.05, 2.7, 2.6, .7, 31, TEAL, True)
text(s, '200 / class', 10.05, 3.85, 2.6, .45, 20, MUTED)
text(s, f"{float(VALUES[200]['mean']):.6f}", 10.05, 4.45, 2.6, .7, 31, TEAL, True)
text(s, 'SD is episode spread, not a confidence interval.', .85, 6.55, 11.6, .32, 16, MUTED)

s = slide('A strong fixed-city result—not universal superiority', '06 · Result interpretation',
          'Paired evidence: candidate/reports/RESULTS.md and STATISTICAL_ANALYSIS.md',
          'Speaker 3. At 200/class, compare Coordinate-RF with ASTRA on matched episodes. Do not infer independent-city significance from the paired tests.')
card(s, .7, 2, 5.8, '200 labels/class', 'Coordinate-RF: 0.749299 ± 0.005716\nMatched ASTRA: 0.733037\nMean gain: +0.016262', 3.15)
card(s, 6.8, 2, 5.8, 'What the evidence says', '200/200 paired wins at 200/class\nFive-shot superiority is inconclusive\nOne shared target-city population', 3.15)
text(s, 'The 800-label development bank and current support are excluded from final queries.', .85, 5.8, 11.7, .85, 23, MUTED)

s = slide('Spatial extrapolation reveals the trade-off', '07 · Analysis & limitations',
          'Frozen spatial evidence: candidate/reports/SPATIAL_ANALYSIS.md',
          'Speaker 3 → Speaker 4. Lower spatial means are material negative evidence. Different query geography prevents equating spatial and random-pixel scores.')
card(s, .7, 2, 5.8, 'Four-direction mean', 'Coordinate-RF: 0.694216\nMatched ASTRA: 0.696672', 2.65)
card(s, 6.8, 2, 5.8, 'Worst-direction mean', 'Coordinate-RF: 0.644816\nMatched ASTRA: 0.658183', 2.65)
text(s, 'Coordinates help within this pool; they do not establish transfer to new geography.', .85, 5.25, 11.6, 1.1, 28, ORANGE, True)

s = slide('Frozen evidence, explicit reproducibility boundaries', '08 · Engineering & integrity',
          'Verification: scripts/validate_release.py; FINAL_SUBMISSION/manifest.json',
          'Speaker 4. Hash checks are safe; --verify is not just a checksum command and fits support RFs. Explain the pool-bound artifact and distinguish static checks from runtime reproduction.')
card(s, .7, 2, 5.8, 'What is preserved', 'Model + configuration + ordered inputs\nNotebook outputs + result tables\nManifest and SHA-256 checksums', 3.35)
card(s, 6.8, 2, 5.8, 'What is not claimed', 'Fresh training or runtime replay here\nA self-contained raw-data rebuild\nOrganizer-held-out validation', 3.35)
text(s, 'Safe audit: python -B scripts/validate_release.py', .85, 5.9, 11.6, .65, 24, TEAL, True)

s = slide('A clear result—and a clear next test', '09 · Conclusion & future work',
          'Start here: docs/judge/START_HERE.md · results remain locked',
          'Speaker 4. Close with exact score and scope. Future work is proposed, not performed. All four speakers should rehearse handoffs. Take questions on source transfer, the 25-shot regime, extra supervision and spatial limitations.')
text(s, 'Coordinate-RF · macro-F1 0.749299', .75, 2, 11.85, .85, 38, TEAL, True)
text(s, '200 labels/class · fixed-city random-pixel evaluation', .78, 3, 11.7, .65, 25)
card(s, .7, 4, 5.8, 'Before competition upload', 'Confirm full-pool, coordinate and\nadditional-supervision eligibility.', 2.35)
card(s, 6.8, 4, 5.8, 'Next research—not run', 'Independent-city / spatial validation\nPortable inference for new target pools', 2.35)

assert len(prs.slides) == 10
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        assert min(sh.left, sh.top) >= 0, (i, sh.name)
        assert sh.left + sh.width <= prs.slide_width + 2, (i, sh.name)
        assert sh.top + sh.height <= prs.slide_height + 2, (i, sh.name)
prs.save(OUT/'Coordinate-RF_Judges.pptx')
receipt = {'slides': 10, 'scientific_execution': False,
           'source_table': TABLE.relative_to(ROOT).as_posix(),
           'source_sha256': hashlib.sha256(TABLE.read_bytes()).hexdigest(),
           'model_display_name': 'Coordinate-RF', 'locked_macro_f1': '0.749299',
           'figure_policy': 'New presentation view of unchanged CSV means/SD; canonical figures untouched.'}
(ROOT/'reproducibility/release/deck_build.json').write_text(json.dumps(receipt, indent=2))
print('Built 10 slides; no model imports, fitting or prediction generation.')
